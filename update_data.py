import sys
import subprocess

# 1. Verificar e instalar dependencias antes de importar el resto
REQUIRED_LIBS = ['streamlit', 'pandas', 'openpyxl', 'xlrd', 'xlsxwriter']
missing_libs = []
for lib in REQUIRED_LIBS:
    try:
        __import__(lib)
    except ImportError:
        missing_libs.append(lib)

if missing_libs:
    print(f"Librerías faltantes detectadas: {missing_libs}")
    print("Instalando dependencias necesarias. Por favor espera...")
    try:
        # Ejecutar pip install usando el ejecutable actual de python
        subprocess.check_call([sys.executable, "-m", "pip", "install"] + missing_libs)
        print("¡Dependencias instaladas con éxito!")
    except Exception as e:
        print(f"Error crítico al instalar dependencias con pip: {e}")
        sys.exit(1)

# 2. Importaciones estándar ahora que sabemos que están instaladas
import os
import glob
import re
import pandas as pd
import xlrd
import openpyxl
from datetime import datetime

# Directores de trabajo
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(BASE_DIR, "datos_consolidados.csv")

def clean_header_name(h):
    """Limpia el nombre de la columna para facilitar el mapeo."""
    if not h:
        return ""
    h_clean = " ".join(str(h).split()).upper().strip()
    return h_clean

def read_xls_file(file_path):
    """Lee archivos .xls antiguos usando xlrd, limpia y devuelve lista de DataFrames por solapa."""
    try:
        book = xlrd.open_workbook(file_path)
    except Exception as e:
        print(f"Error al abrir {file_path} con xlrd: {e}")
        return []

    sheet_dfs = []
    for sheet_name in book.sheet_names():
        sheet = book.sheet_by_name(sheet_name)
        if sheet.nrows <= 1:
            continue
        
        headers = [clean_header_name(sheet.cell_value(0, col)) for col in range(sheet.ncols)]
        sheet_rows = []
        for r in range(1, sheet.nrows):
            row_dict = {
                '_proveedor': sheet_name.strip(),
                '_archivo': os.path.basename(file_path)
            }
            for c in range(min(len(headers), sheet.ncols)):
                col_name = headers[c]
                if not col_name:
                    continue
                
                val = sheet.cell_value(r, c)
                ctype = sheet.cell_type(r, c)
                
                if ctype == xlrd.XL_CELL_DATE:
                    try:
                        date_tuple = xlrd.xldate_as_tuple(val, book.datemode)
                        val = datetime(*date_tuple)
                    except:
                        pass
                elif ctype == xlrd.XL_CELL_EMPTY or val == '':
                    val = None
                
                row_dict[col_name] = val
            
            if any(v is not None for k, v in row_dict.items() if k not in ['_proveedor', '_archivo']):
                sheet_rows.append(row_dict)
                
        if sheet_rows:
            sheet_df = pd.DataFrame(sheet_rows)
            cleaned_df = map_and_clean_df(sheet_df)
            if not cleaned_df.empty:
                sheet_dfs.append(cleaned_df)
    return sheet_dfs

def read_xlsx_file(file_path):
    """Lee archivos .xlsx nuevos usando openpyxl, limpia y devuelve lista de DataFrames por solapa."""
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        print(f"Error al abrir {file_path} con openpyxl: {e}")
        return []

    sheet_dfs = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        if sheet.max_row <= 1:
            continue
        
        headers = []
        for cell in sheet[1]:
            headers.append(clean_header_name(cell.value))
            
        sheet_rows = []
        for r in range(2, sheet.max_row + 1):
            row_dict = {
                '_proveedor': sheet_name.strip(),
                '_archivo': os.path.basename(file_path)
            }
            for c in range(1, len(headers) + 1):
                col_name = headers[c-1]
                if not col_name:
                    continue
                cell = sheet.cell(row=r, column=c)
                val = cell.value
                if val == '':
                    val = None
                row_dict[col_name] = val
                
            if any(v is not None for k, v in row_dict.items() if k not in ['_proveedor', '_archivo']):
                sheet_rows.append(row_dict)
                
        if sheet_rows:
            sheet_df = pd.DataFrame(sheet_rows)
            cleaned_df = map_and_clean_df(sheet_df)
            if not cleaned_df.empty:
                sheet_dfs.append(cleaned_df)
    return sheet_dfs

def map_and_clean_df(df):
    """Asocia columnas encontradas a las columnas estándar requeridas."""
    if df.empty:
        return pd.DataFrame()
        
    mapped_df = pd.DataFrame()
    cols = df.columns
    
    # 1. FECHA
    col_fecha = [c for c in cols if 'FECHA' in c or 'DATE' in c]
    if col_fecha:
        mapped_df['fecha'] = df[col_fecha[0]]
    else:
        mapped_df['fecha'] = pd.Series([None]*len(df))
        
    # 2. FACTURA (Busca por expresiones regulares para soportar FC, FC Nº, FC. Nº, Factura, etc.)
    import re
    col_fc = []
    for c in cols:
        if re.search(r'\bFC\b|\bFC\W|FACTURA|INVOICE', c):
            col_fc.append(c)
            break
            
    if col_fc:
        mapped_df['nro_fc'] = df[col_fc[0]]
    else:
        mapped_df['nro_fc'] = pd.Series([None]*len(df))
        
    # 3. GUIA/LOTE
    col_guia = [c for c in cols if 'GUIA' in c or 'LOTE' in c or 'TRACKING' in c]
    if col_guia:
        mapped_df['guia'] = df[col_guia[0]]
    else:
        mapped_df['guia'] = pd.Series([None]*len(df))
        
    # 4. ARTICULO
    col_art = [c for c in cols if 'ARTICULO' in c or 'PRODUCTO' in c or 'PRODUCT' in c or 'ITEM' in c]
    if col_art:
        mapped_df['articulo_original'] = df[col_art[0]]
    else:
        mapped_df['articulo_original'] = pd.Series([None]*len(df))
        
    # 5. KILOS UNITARIOS (KG.)
    col_kg = []
    for c in cols:
        # Buscar 'KG' o 'KILO' pero excluir si contiene palabras de costo/pesos
        if ('KG' in c or 'KILO' in c) and not any(k in c for k in ['COSTO', 'PRECIO', 'PESOS', 'PAGAR']):
            col_kg.append(c)
            break
            
    if col_kg:
        mapped_df['kg_unitario'] = df[col_kg[0]]
    else:
        mapped_df['kg_unitario'] = pd.Series([None]*len(df))
        
    # 6. CANTIDAD (CANT)
    col_cant = [c for c in cols if 'CANT' in c or 'QTY' in c or 'CANTIDAD' in c]
    if col_cant:
        mapped_df['cantidad'] = df[col_cant[0]]
    else:
        mapped_df['cantidad'] = pd.Series([None]*len(df))
        
    # 7. COSTO FINAL EN PESOS (Unitario)
    col_costo = []
    for c in cols:
        if 'COSTO FINAL' in c and 'KILO' not in c and 'PESOS' in c:
            col_costo.append(c)
            break
    if not col_costo:
        for c in cols:
            if 'COSTO' in c and 'PESOS' in c and 'KILO' not in c:
                col_costo.append(c)
                break
                
    if col_costo:
        mapped_df['costo_final_unitario'] = df[col_costo[0]]
    else:
        mapped_df['costo_final_unitario'] = pd.Series([None]*len(df))
        
    # 8. COSTO FINAL POR KILO (Buscar específicamente el costo final en pesos por kilo)
    col_costo_kilo = []
    # Intento 1: Buscar columna que contenga COSTO, PESOS y KILO (e.g. COSTO FINAL EN PESOS POR KILO)
    for c in cols:
        if 'KILO' in c and 'COSTO' in c and 'PESOS' in c:
            col_costo_kilo.append(c)
            break
            
    # Intento 2: Si no se encuentra, buscar COSTO y KILO (evitando USD, FOB, FLETE)
    if not col_costo_kilo:
        for c in cols:
            if 'KILO' in c and 'COSTO' in c and not any(x in c for x in ['USD', 'FOB', 'FLETE']):
                col_costo_kilo.append(c)
                break
                
    # Intento 3: Si sigue sin encontrarse, buscar cualquier columna con KILO y (COSTO o PRECIO o VALOR)
    if not col_costo_kilo:
        for c in cols:
            if 'KILO' in c and ('COSTO' in c or 'PRECIO' in c or 'VALOR' in c):
                col_costo_kilo.append(c)
                break
                
    if col_costo_kilo:
        mapped_df['costo_por_kilo'] = df[col_costo_kilo[0]]
    else:
        mapped_df['costo_por_kilo'] = pd.Series([None]*len(df))

    # Conservamos columnas de metadatos
    mapped_df['proveedor'] = df['_proveedor'] if '_proveedor' in df.columns else 'DESCONOCIDO'
    mapped_df['archivo_origen'] = df['_archivo'] if '_archivo' in df.columns else ''
    
    # Rellenar valores agrupados (Forward Fill)
    mapped_df['fecha'] = mapped_df['fecha'].ffill()
    mapped_df['nro_fc'] = mapped_df['nro_fc'].ffill()
    
    # Procesamiento y limpieza de datos
    mapped_df['fecha'] = pd.to_datetime(mapped_df['fecha'], errors='coerce')
    
    # Convertir números
    mapped_df['kg_unitario'] = pd.to_numeric(mapped_df['kg_unitario'], errors='coerce').fillna(0.0)
    mapped_df['cantidad'] = pd.to_numeric(mapped_df['cantidad'], errors='coerce').fillna(0.0)
    mapped_df['costo_final_unitario'] = pd.to_numeric(mapped_df['costo_final_unitario'], errors='coerce').fillna(0.0)
    mapped_df['costo_por_kilo'] = pd.to_numeric(mapped_df['costo_por_kilo'], errors='coerce').fillna(0.0)
    
    # Convertir guía y fc a strings legibles
    mapped_df['guia'] = mapped_df['guia'].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
    mapped_df.loc[mapped_df['guia'].isin(['nan', 'None', '']), 'guia'] = '-'
    
    mapped_df['nro_fc'] = mapped_df['nro_fc'].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
    mapped_df.loc[mapped_df['nro_fc'].isin(['nan', 'None', '']), 'nro_fc'] = '-'
    
    # Separar Artículo en Producto y Variedad
    def parse_article(desc):
        if not desc or pd.isna(desc):
            return "S/D", "S/D"
        desc = " ".join(str(desc).split()).strip().upper()
        parts = desc.split(maxsplit=1)
        if len(parts) == 2:
            return parts[0], parts[1]
        return parts[0], "SIN ESPECIFICAR"
        
    articles_split = mapped_df['articulo_original'].apply(parse_article)
    mapped_df['articulo_producto'] = [x[0] for x in articles_split]
    mapped_df['articulo_variedad'] = [x[1] for x in articles_split]
    
    # Calcular totales
    mapped_df['kilos_totales'] = mapped_df['kg_unitario'] * mapped_df['cantidad']
    mapped_df['costo_final_total'] = mapped_df['costo_final_unitario'] * mapped_df['cantidad']
    
    # Si el costo por kilo vino vacío o en cero, calcularlo
    mask_cero_kilo = (mapped_df['costo_por_kilo'] == 0.0) & (mapped_df['kg_unitario'] > 0.0)
    mapped_df.loc[mask_cero_kilo, 'costo_por_kilo'] = mapped_df['costo_final_unitario'] / mapped_df['kg_unitario']
    
    mapped_df = mapped_df.dropna(subset=['articulo_original'])
    
    return mapped_df

def consolidate_data():
    print(f"Buscando archivos excel en: {BASE_DIR}")
    
    excel_files = []
    for ext in ['*.xls', '*.xlsx']:
        for f in glob.glob(os.path.join(BASE_DIR, ext)):
            filename = os.path.basename(f)
            if not filename.startswith("~$") and "datos_consolidados" not in filename and "dashboard_export" not in filename:
                excel_files.append(f)
                
    print(f"Archivos encontrados: {[os.path.basename(f) for f in excel_files]}")
    
    all_dfs = []
    for f in excel_files:
        print(f"Procesando archivo: {os.path.basename(f)}")
        if f.endswith('.xls'):
            dfs = read_xls_file(f)
        else:
            dfs = read_xlsx_file(f)
        all_dfs.extend(dfs)
        
    if not all_dfs:
        print("No se encontraron datos en los archivos de Excel.")
        return
        
    df_clean = pd.concat(all_dfs, ignore_index=True)
    df_clean = df_clean.sort_values(by='fecha', ascending=False)
    df_clean.to_csv(OUTPUT_FILE, index=False, encoding='utf-8-sig')
    print(f"Datos consolidados guardados en: {OUTPUT_FILE} ({len(df_clean)} registros)")

if __name__ == "__main__":
    consolidate_data()
